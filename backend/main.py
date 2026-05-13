import csv
import io
import json
import os
import re
import subprocess
import tempfile
import uuid
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import FastAPI, UploadFile, File, Form, Cookie
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from fastapi.staticfiles import StaticFiles

from database import init_db, get_db
from stop_words import STOP_WORDS, STOP_WORDS_RU, STOP_WORDS_EN

CPP_EXECUTABLE = os.path.join(os.path.dirname(__file__), "analyzer")
DATA_DIR = os.path.join(os.path.dirname(__file__), "data")
STOP_WORDS_FILE = os.path.join(DATA_DIR, "stop_words.txt")
STOP_WORDS_FILES = {
    "ru": os.path.join(DATA_DIR, "stop_words_ru.txt"),
    "en": os.path.join(DATA_DIR, "stop_words_en.txt"),
}
MAX_UPLOAD_BYTES = 50 * 1024 * 1024  # 50 MB
ALLOWED_EXTENSIONS = {".txt", ".docx", ".pdf"}


class ExtractionError(Exception):
    """Raised by extract_text() to signal a user-facing failure."""

    def __init__(self, status_code: int, message: str):
        self.status_code = status_code
        self.message = message
        super().__init__(message)


def extract_text(filename: str, content: bytes) -> str:
    """Extract plain text from a .txt/.docx/.pdf upload.

    Raises ExtractionError(status_code, message) on user-facing failures so the
    /api/analyze handler can map them to the right HTTP response.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        # Prefer charset-normalizer if available — handles UTF-8/16, CP1251, etc.
        try:
            from charset_normalizer import from_bytes
            best = from_bytes(content).best()
            if best is not None:
                return str(best)
        except ImportError:
            pass
        return content.decode("utf-8", errors="replace")

    if ext == ".docx":
        try:
            import docx  # python-docx
        except ImportError:
            raise ExtractionError(501, "python-docx не установлен на сервере")
        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception:
            raise ExtractionError(422, "Не удалось прочитать .docx файл (повреждён?)")
        return "\n".join(p.text for p in doc.paragraphs)

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ExtractionError(501, "pypdf не установлен на сервере")
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ExtractionError(422, "PDF зашифрован — расшифровка не поддерживается")
            parts = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(parts)
        except ExtractionError:
            raise
        except Exception:
            raise ExtractionError(422, "Не удалось прочитать .pdf файл (повреждён?)")

    raise ExtractionError(
        400,
        f"Неподдерживаемый формат: {ext or 'unknown'}. Используйте .txt, .docx или .pdf",
    )


@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    yield


app = FastAPI(lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_or_create_session(session_id: str | None) -> str:
    if session_id:
        with get_db() as db:
            row = db.execute("SELECT id FROM sessions WHERE id = ?", (session_id,)).fetchone()
            if row:
                return session_id
    new_id = uuid.uuid4().hex
    with get_db() as db:
        db.execute("INSERT INTO sessions (id) VALUES (?)", (new_id,))
    return new_id


@app.post("/api/analyze")
async def analyze(
    file: UploadFile = File(...),
    exclude_stopwords: bool = Form(default=False),
    language: str = Form(default="auto"),
    min_length: int = Form(default=0),
    top_n: int = Form(default=0),
    ignore_numbers: bool = Form(default=False),
    ngram: int = Form(default=1),
    normalize: str = Form(default="none"),
    session_id: str | None = Cookie(default=None),
):
    if not file.filename:
        return JSONResponse(status_code=400, content={"error": "Имя файла отсутствует"})

    content = await file.read()
    if len(content) > MAX_UPLOAD_BYTES:
        return JSONResponse(
            status_code=413,
            content={"error": f"Файл слишком большой (максимум {MAX_UPLOAD_BYTES // (1024 * 1024)} МБ)"},
        )

    try:
        text = extract_text(file.filename, content)
    except ExtractionError as e:
        return JSONResponse(status_code=e.status_code, content={"error": e.message})

    # Clamp inputs so a malicious/typo'd value can't reach the C++ CLI as-is.
    if ngram not in (1, 2, 3):
        ngram = 1
    if normalize not in ("none", "stem"):
        normalize = "none"

    sid = get_or_create_session(session_id)

    with tempfile.NamedTemporaryFile(delete=False, suffix=".txt", mode="w", encoding="utf-8") as tmp:
        tmp.write(text)
        tmp_path = tmp.name

    try:
        result = run_analyzer(
            tmp_path,
            exclude_stopwords=exclude_stopwords,
            language=language,
            ngram=ngram,
            normalize=normalize,
        )
    finally:
        os.unlink(tmp_path)

    result = post_process(
        result,
        min_length=min_length,
        top_n=top_n,
        ignore_numbers=ignore_numbers,
    )

    with get_db() as db:
        cur = db.execute(
            "INSERT INTO analysis_history "
            "(session_id, filename, result_json, total_words, unique_words) "
            "VALUES (?, ?, ?, ?, ?)",
            (
                sid,
                file.filename,
                json.dumps(result),
                int(result.get("total_words") or 0),
                int(result.get("unique_words") or 0),
            ),
        )
        entry_id = cur.lastrowid

    response = JSONResponse(content={**result, "id": entry_id})
    response.set_cookie("session_id", sid, httponly=True, samesite="lax")
    return response


def post_process(
    result: dict,
    min_length: int,
    top_n: int,
    ignore_numbers: bool = False,
) -> dict:
    """Apply min_length / ignore_numbers filters and top_n truncation.

    Filters are applied to `frequencies` (unigrams). For `ngrams` only top_n
    truncation is applied — min_length doesn't make sense for phrases and
    ignore_numbers would need per-token inspection inside the phrase.
    """
    freqs = result.get("frequencies", {})
    if ignore_numbers:
        freqs = {w: c for w, c in freqs.items() if not w.isdigit()}
    if min_length > 0:
        freqs = {w: c for w, c in freqs.items() if len(w) >= min_length}
    if top_n > 0:
        freqs = dict(list(freqs.items())[:top_n])

    out = {**result, "frequencies": freqs}

    ngrams = result.get("ngrams")
    if isinstance(ngrams, dict) and top_n > 0:
        out["ngrams"] = {
            n: dict(list(phrases.items())[:top_n])
            for n, phrases in ngrams.items()
        }
    return out


def _resolve_stopwords_path(language: str) -> str:
    """Pick the bundled stop-words file for the requested language; fall back
    to the combined list when language is 'auto' or unrecognized."""
    path = STOP_WORDS_FILES.get(language)
    if path and os.path.isfile(path):
        return path
    return STOP_WORDS_FILE


def run_analyzer(
    file_path: str,
    exclude_stopwords: bool,
    language: str = "auto",
    ngram: int = 1,
    normalize: str = "none",
) -> dict:
    exe = CPP_EXECUTABLE
    if os.name == "nt":
        exe += ".exe"

    if not os.path.isfile(exe):
        return fallback_analyze(
            file_path,
            exclude_stopwords=exclude_stopwords,
            language=language,
            ngram=ngram,
        )

    cmd = [exe]
    if exclude_stopwords:
        cmd += ["--stopwords", _resolve_stopwords_path(language)]
    if ngram >= 2:
        cmd += ["--ngram", str(ngram)]
    if normalize == "stem":
        cmd += ["--normalize", "stem"]
    cmd.append(file_path)

    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
    if proc.returncode != 0:
        return fallback_analyze(
            file_path,
            exclude_stopwords=exclude_stopwords,
            language=language,
            ngram=ngram,
        )

    return json.loads(proc.stdout)


# Mirrors the C++ tokenizer: ASCII alnum + Cyrillic letters, lowercased, any
# other codepoint (punctuation, whitespace, unicode dashes) acts as a separator.
_TOKEN_RE = re.compile(r"[a-z0-9Ѐ-ӿ]+")


def fallback_analyze(
    file_path: str,
    exclude_stopwords: bool,
    language: str = "auto",
    ngram: int = 1,
) -> dict:
    """Python fallback used when the C++ analyzer binary is unavailable.

    Implements stop-word filtering and n-gram extraction so the fallback path
    stays feature-comparable to the C++ binary for the parts that don't need
    Porter stemming. `normalize` is accepted by the caller but not honored
    here — Python-side stemming would defeat the purpose of having a fast
    fallback, and tests live in C++.
    """
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()

    tokens = _TOKEN_RE.findall(text.lower())
    total = len(tokens)

    if language == "ru":
        sw = STOP_WORDS_RU
    elif language == "en":
        sw = STOP_WORDS_EN
    else:
        sw = STOP_WORDS

    freq: dict[str, int] = {}
    kept: list[str] = []
    for token in tokens:
        if exclude_stopwords and token in sw:
            continue
        freq[token] = freq.get(token, 0) + 1
        kept.append(token)

    sorted_freq = dict(sorted(freq.items(), key=lambda x: x[1], reverse=True))

    # Mirror C++ detect_language: count Cyrillic vs ASCII letters across all
    # tokens; majority wins. "unknown" when no letters seen at all.
    cyr = lat = 0
    for tok in tokens:
        for ch in tok:
            if "a" <= ch <= "z":
                lat += 1
            elif "\u0400" <= ch <= "\u04FF":
                cyr += 1
    if cyr == 0 and lat == 0:
        detected_lang = "unknown"
    else:
        detected_lang = "ru" if cyr >= lat else "en"

    out = {
        "total_words": total,
        "unique_words": len(freq),
        "frequencies": sorted_freq,
        "language": detected_lang,
    }

    if ngram >= 2 and len(kept) >= ngram:
        ng: dict[str, int] = {}
        for i in range(len(kept) - ngram + 1):
            phrase = " ".join(kept[i:i + ngram])
            ng[phrase] = ng.get(phrase, 0) + 1
        ng_sorted = dict(sorted(ng.items(), key=lambda x: x[1], reverse=True))
        out["ngrams"] = {str(ngram): ng_sorted}

    return out


@app.get("/api/history")
async def history(session_id: str | None = Cookie(default=None)):
    if not session_id:
        return []

    # Summary-only: avoids parsing the (potentially huge) result_json blob for
    # every entry. The frontend list UI only renders total_words; the full
    # result is fetched lazily via GET /api/history/{id} when an entry is
    # opened or used in a comparison.
    with get_db() as db:
        rows = db.execute(
            "SELECT id, filename, total_words, unique_words, created_at "
            "FROM analysis_history "
            "WHERE session_id = ? ORDER BY created_at DESC",
            (session_id,),
        ).fetchall()

    return [
        {
            "id": row["id"],
            "filename": row["filename"],
            "total_words": row["total_words"],
            "unique_words": row["unique_words"],
            "created_at": row["created_at"],
        }
        for row in rows
    ]


@app.get("/api/history/{entry_id}")
async def history_entry(entry_id: int, session_id: str | None = Cookie(default=None)):
    if not session_id:
        return JSONResponse(status_code=401, content={"error": "No session"})

    with get_db() as db:
        row = db.execute(
            "SELECT id, filename, result_json, created_at FROM analysis_history "
            "WHERE id = ? AND session_id = ?",
            (entry_id, session_id),
        ).fetchone()

    if not row:
        return JSONResponse(status_code=404, content={"error": "Entry not found"})

    return {
        "id": row["id"],
        "filename": row["filename"],
        "result": json.loads(row["result_json"]),
        "created_at": row["created_at"],
    }


@app.delete("/api/history/{entry_id}")
async def delete_history_entry(entry_id: int, session_id: str | None = Cookie(default=None)):
    if not session_id:
        return JSONResponse(status_code=401, content={"error": "No session"})

    with get_db() as db:
        db.execute(
            "DELETE FROM analysis_history WHERE id = ? AND session_id = ?",
            (entry_id, session_id),
        )

    return {"ok": True}


@app.get("/api/export/{entry_id}")
async def export_entry(
    entry_id: int,
    format: str = "csv",
    session_id: str | None = Cookie(default=None),
):
    if not session_id:
        return JSONResponse(status_code=401, content={"error": "No session"})

    with get_db() as db:
        row = db.execute(
            "SELECT filename, result_json FROM analysis_history "
            "WHERE id = ? AND session_id = ?",
            (entry_id, session_id),
        ).fetchone()

    if not row:
        return JSONResponse(status_code=404, content={"error": "Entry not found"})

    result = json.loads(row["result_json"])
    base_name = os.path.splitext(row["filename"])[0] or "export"
    freqs = result.get("frequencies", {})

    if format == "json":
        return Response(
            content=json.dumps(result, ensure_ascii=False, indent=2),
            media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="{base_name}.json"'},
        )

    if format == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(["word", "count"])
        for word, count in freqs.items():
            writer.writerow([word, count])
        # BOM so Excel auto-detects UTF-8.
        content = "\ufeff" + buf.getvalue()
        return Response(
            content=content,
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{base_name}.csv"'},
        )

    if format == "xlsx":
        try:
            from openpyxl import Workbook
        except ImportError:
            return JSONResponse(
                status_code=501,
                content={"error": "openpyxl не установлен на сервере"},
            )
        wb = Workbook()
        ws = wb.active
        ws.title = "Frequencies"
        ws.append(["word", "count"])
        for word, count in freqs.items():
            ws.append([word, count])
        out = io.BytesIO()
        wb.save(out)
        return Response(
            content=out.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{base_name}.xlsx"'},
        )

    return JSONResponse(
        status_code=400,
        content={"error": f"Unknown format: {format!r}. Use csv, xlsx, or json."},
    )


# ─── Static frontend serving (combined Docker deploy) ───────────────
# When the built frontend exists at ./static (see root Dockerfile), mount it
# so the single service handles both API and SPA without nginx.
_static_dir = Path(__file__).parent / "static"
if _static_dir.is_dir():
    from fastapi.responses import FileResponse

    # Serve static assets (JS, CSS, images, etc.)
    app.mount("/assets", StaticFiles(directory=_static_dir / "assets"), name="assets")

    # Serve other static files at root (favicon, icons, etc.)
    @app.get("/{full_path:path}")
    async def spa_fallback(full_path: str):
        """Serve static file if it exists, otherwise index.html (SPA routing)."""
        file_path = _static_dir / full_path
        if full_path and file_path.is_file():
            return FileResponse(file_path)
        return FileResponse(_static_dir / "index.html")
